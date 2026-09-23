from pathlib import Path

import pytest
from conftest import DEMO_PDF

from app.pdf import convert
from app.pdf import extract as x
from app.video.ffmpeg import concat_audio, encode_video, media_duration
from app.tts.synth import MockTTS

SECTIONS = [("1 Networks", "A network connects hosts. " * 20), ("2 TCP", "TCP is reliable and ordered. " * 20)]


def _make(tmp: Path, kind: str) -> Path:
    if kind == "md":
        p = tmp / "n.md"
        p.write_text("# Notes\n\n" + "\n\n".join(f"## {t}\n\n{b}" for t, b in SECTIONS), encoding="utf-8")
    elif kind == "txt":
        p = tmp / "n.txt"
        p.write_text("Notes\n\n" + "\n\n".join(f"{t}\n\n{b}" for t, b in SECTIONS), encoding="utf-8")
    elif kind == "docx":
        import docx

        p = tmp / "n.docx"
        d = docx.Document()
        d.add_heading("Notes", 0)
        for t, b in SECTIONS:
            d.add_heading(t, 1)
            d.add_paragraph(b)
        d.save(p)
    else:
        from pptx import Presentation

        p = tmp / "n.pptx"
        prs = Presentation()
        for t, b in SECTIONS:
            s = prs.slides.add_slide(prs.slide_layouts[1])
            s.shapes.title.text = t
            s.placeholders[1].text_frame.text = b[:200]
        prs.save(p)
    return p


@pytest.mark.parametrize("kind", ["md", "txt", "docx", "pptx"])
def test_convert_keeps_structure(tmp_path, kind):
    src = _make(tmp_path, kind)
    pdf = convert.convert_to_pdf(src, src.name, tmp_path / "out.pdf")
    ex = x.extract(pdf)
    heads = " ".join(h for p in ex.pages for h in p.headings)
    assert ex.page_count >= 1
    assert "Networks" in heads and "TCP" in heads


def test_extract_demo_pdf_and_suitability():
    ex = x.extract(DEMO_PDF)
    assert ex.page_count == 10
    assert x.suitability(ex).quality == "good"


def test_ffmpeg_encode_roundtrip(tmp_path):
    """Regression for the Linux-only 'flush of closed file' pipe bug."""
    mp3 = tmp_path / "a.mp3"
    n = MockTTS().synthesize("Hello there, this is a short test.", mp3)
    audio = tmp_path / "a.m4a"
    concat_audio([(mp3, 0.5)], audio)
    frames = (bytes([255, 255, 255]) * (64 * 64) for _ in range(int((n.duration + 0.5) * 8)))
    out = tmp_path / "v.mp4"
    encode_video(frames, 64, 64, 8, audio, out)
    assert out.exists() and out.stat().st_size > 1000
    assert media_duration(out) > 1
