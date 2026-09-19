"""Convert non-PDF study material (txt, md, docx, pptx) into a clean PDF so the rest of
the pipeline - page numbers, page images for the brain, source references - is unchanged.
Headings are preserved as their own lines so the heading detector still finds structure."""
from __future__ import annotations

import os
import re
from pathlib import Path

from fpdf import FPDF

SUPPORTED_EXT = {".pdf", ".txt", ".md", ".markdown", ".docx", ".pptx"}
ACCEPT_LABEL = "PDF, DOCX, PPTX, TXT or Markdown"

Block = tuple[str, str]   # (kind, text): kind in title|h1|h2|h3|p|li|code|pagebreak

_FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    r"C:\Windows\Fonts\arial.ttf",
    r"C:\Windows\Fonts\segoeui.ttf",
    "/System/Library/Fonts/Supplemental/Arial.ttf",
]


def source_type(filename: str) -> str:
    return Path(filename or "").suffix.lower().lstrip(".") or "pdf"


def is_supported(filename: str) -> bool:
    return Path(filename or "").suffix.lower() in SUPPORTED_EXT


# ------------------------------------------------------------------ readers
_MD_INLINE = [
    (re.compile(r"!\[([^\]]*)\]\([^)]*\)"), r"\1"),      # images -> alt
    (re.compile(r"\[([^\]]+)\]\([^)]*\)"), r"\1"),       # links -> text
    (re.compile(r"(\*\*|__)(.+?)\1"), r"\2"),            # bold
    (re.compile(r"(\*|_)(.+?)\1"), r"\2"),               # italic
    (re.compile(r"`([^`]+)`"), r"\1"),                   # inline code
]


def _clean_inline(s: str) -> str:
    for rx, rep in _MD_INLINE:
        s = rx.sub(rep, s)
    return s.strip()


def read_markdown(text: str) -> list[Block]:
    blocks: list[Block] = []
    para: list[str] = []
    in_code = False

    def flush():
        if para:
            blocks.append(("p", " ".join(para)))
            para.clear()

    for raw in text.splitlines():
        line = raw.rstrip()
        if line.strip().startswith("```"):
            flush()
            in_code = not in_code
            continue
        if in_code:
            blocks.append(("code", line))
            continue
        if not line.strip():
            flush()
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", line)
        if m:
            flush()
            level = min(3, len(m.group(1)))
            blocks.append((f"h{level}", _clean_inline(m.group(2).strip("# "))))
            continue
        if re.match(r"^\s*([-*+]|\d+[.)])\s+", line):
            flush()
            blocks.append(("li", _clean_inline(re.sub(r"^\s*([-*+]|\d+[.)])\s+", "", line))))
            continue
        if line.startswith("---") or line.startswith("==="):
            flush()
            continue
        para.append(_clean_inline(line))
    flush()
    return blocks


def read_text(text: str) -> list[Block]:
    """Plain text: blank lines separate paragraphs; short Title-Case lines become headings."""
    blocks: list[Block] = []
    para: list[str] = []
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            if para:
                blocks.append(("p", " ".join(para)))
                para.clear()
            continue
        words = line.split()
        if (len(line) < 70 and not line.endswith((".", ",", ";")) and len(words) <= 10
                and sum(w[:1].isupper() for w in words) >= max(1, len(words) * 0.6)):
            if para:
                blocks.append(("p", " ".join(para)))
                para.clear()
            blocks.append(("h2", line))
        else:
            para.append(line)
    if para:
        blocks.append(("p", " ".join(para)))
    return blocks


def read_docx(path: Path) -> list[Block]:
    import docx  # python-docx

    d = docx.Document(str(path))
    blocks: list[Block] = []
    for p in d.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style.startswith("title"):
            blocks.append(("title", text))
        elif style.startswith("heading"):
            m = re.search(r"(\d+)", style)
            level = min(3, int(m.group(1))) if m else 2
            blocks.append((f"h{level}", text))
        elif "list" in style:
            blocks.append(("li", text))
        else:
            blocks.append(("p", text))
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                blocks.append(("li", " | ".join(cells)))
    return blocks


def read_pptx(path: Path) -> list[Block]:
    from pptx import Presentation  # python-pptx

    prs = Presentation(str(path))
    blocks: list[Block] = []
    for i, slide in enumerate(prs.slides, 1):
        if i > 1:
            blocks.append(("pagebreak", ""))
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip()
        blocks.append(("h1", title or f"Slide {i}"))
        for shape in slide.shapes:
            if shape == slide.shapes.title or not getattr(shape, "has_text_frame", False):
                continue
            for para in shape.text_frame.paragraphs:
                text = "".join(r.text for r in para.runs).strip()
                if text:
                    blocks.append(("li" if para.level else "p", text))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                blocks.append(("p", f"Notes: {notes}"))
    return blocks


# ------------------------------------------------------------------ writer
class _Doc(FPDF):
    def __init__(self):
        super().__init__(format="A4")
        self.set_auto_page_break(auto=True, margin=18)
        self.set_margins(18, 18, 18)
        self.unicode = False
        for cand in _FONT_CANDIDATES:
            if os.path.exists(cand):
                try:
                    self.add_font("body", "", cand)
                    self.unicode = True
                    break
                except Exception:  # noqa: BLE001
                    continue
        self.family = "body" if self.unicode else "Helvetica"

    def text_ok(self, s: str) -> str:
        if self.unicode:
            return s
        return s.encode("latin-1", "replace").decode("latin-1")

    def write_block(self, kind: str, text: str) -> None:
        text = self.text_ok(text)
        w = self.w - self.l_margin - self.r_margin
        if kind == "pagebreak":
            self.add_page()
            return
        if kind == "title":
            self.set_font(self.family, size=20)
            self.multi_cell(w, 10, text)
            self.ln(3)
        elif kind == "h1":
            self.set_font(self.family, size=16)
            self.ln(2)
            self.multi_cell(w, 8, text)
            self.ln(1)
        elif kind == "h2":
            self.set_font(self.family, size=13.5)
            self.ln(2)
            self.multi_cell(w, 7, text)
            self.ln(1)
        elif kind == "h3":
            self.set_font(self.family, size=12)
            self.ln(1)
            self.multi_cell(w, 6.5, text)
        elif kind == "li":
            self.set_font(self.family, size=11)
            self.multi_cell(w, 6, "  \u2022 " + text if self.unicode else "  - " + text)
        elif kind == "code":
            self.set_font("Courier", size=9.5)
            self.multi_cell(w, 5, "    " + text)
        else:
            self.set_font(self.family, size=11)
            self.multi_cell(w, 6, text)
            self.ln(2)


def blocks_to_pdf(blocks: list[Block], out_pdf: Path, title: str = "") -> Path:
    pdf = _Doc()
    if title:
        pdf.set_title(title)
    pdf.add_page()
    if not blocks:
        blocks = [("p", "(empty document)")]
    for kind, text in blocks:
        pdf.write_block(kind, text)
    out_pdf.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(out_pdf))
    return out_pdf


def convert_to_pdf(src: Path, filename: str, out_pdf: Path) -> Path:
    """Return a PDF for any supported upload (the file itself when it already is one)."""
    ext = Path(filename).suffix.lower()
    title = Path(filename).stem.replace("_", " ").replace("-", " ").strip()
    if ext == ".pdf":
        return src
    if ext in (".md", ".markdown"):
        text = src.read_text(encoding="utf-8", errors="replace")
        blocks = read_markdown(text)
    elif ext == ".txt":
        text = src.read_text(encoding="utf-8", errors="replace")
        blocks = read_text(text)
    elif ext == ".docx":
        blocks = read_docx(src)
    elif ext == ".pptx":
        blocks = read_pptx(src)
    else:
        raise ValueError(f"Unsupported file type: {ext}")
    if not any(k == "title" or k.startswith("h") for k, _ in blocks[:3]) and title:
        blocks.insert(0, ("title", title))
    return blocks_to_pdf(blocks, out_pdf, title=title)
